from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches, Pt

from .video_outputs import FINAL_VIDEO_REL_PATH, PLACEHOLDER_MP4, ensure_final_video_output


class PptExportError(RuntimeError):
    pass


def export_project_ppt(
    project: dict[str, Any],
    project_dir: Path,
    *,
    allow_placeholder: bool = True,
    sources: dict[str, Any] | None = None,
) -> dict[str, Any]:
    exports_dir = project_dir / "exports"
    exports_dir.mkdir(parents=True, exist_ok=True)

    sources = sources or {}
    pages = _source_pages(sources)
    video_path = _maybe_video_path(project_dir, sources.get("project_config"), allow_placeholder)
    filename = f"{_safe_filename(project.get('name') or project.get('project_id') or 'shan-hai')}-pptx-artifact.pptx"
    pptx_path = exports_dir / filename
    presentation = Presentation()
    notes_count = 0
    if pages:
        for page in pages:
            notes_count += _add_content_slide(presentation, project, page, sources)
    else:
        notes_count += _add_cover_slide(presentation, project)
    if video_path is not None:
        _add_video_slide(presentation, video_path)
    try:
        presentation.save(pptx_path)
    except Exception as exc:
        raise PptExportError(f"PPTX 保存失败：{exc}") from exc

    return {
        "filename": filename,
        "path": f"exports/{filename}",
        "download_url": f"/projects/{project['project_id']}/exports/{filename}",
        "video_path": FINAL_VIDEO_REL_PATH if video_path is not None else None,
        "slide_count": len(presentation.slides),
        "notes_count": notes_count,
        "media_count": _media_count(sources),
    }


def _add_cover_slide(presentation: Presentation, project: dict[str, Any]) -> int:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(8.2), Inches(0.9))
    title_frame = title_box.text_frame
    title_frame.clear()
    title = title_frame.paragraphs[0]
    title.text = project.get("name") or "山海教育公开课"
    title.font.size = Pt(34)
    title.font.bold = True

    meta_box = slide.shapes.add_textbox(Inches(0.95), Inches(2.75), Inches(8), Inches(0.6))
    meta = meta_box.text_frame.paragraphs[0]
    meta.text = f"数学 · {project.get('grade', '')}年级 · {project.get('textbook_version', '')} · {project.get('volume', '')}"
    meta.font.size = Pt(18)
    return 0


def _add_content_slide(presentation: Presentation, project: dict[str, Any], page: dict[str, Any], sources: dict[str, Any]) -> int:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_text = str(page.get("page_objective") or f"第 {page.get('page_index', len(presentation.slides))} 页")
    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(8.8), Inches(0.55))
    title = title_box.text_frame.paragraphs[0]
    title.text = title_text
    title.font.size = Pt(24)
    title.font.bold = True

    visual = page.get("main_visual") if isinstance(page.get("main_visual"), dict) else {}
    left = slide.shapes.add_textbox(Inches(0.7), Inches(1.15), Inches(4.1), Inches(2.1))
    left_frame = left.text_frame
    left_frame.clear()
    left_frame.paragraphs[0].text = str(visual.get("description") or "课堂视觉资产")
    left_frame.paragraphs[0].font.size = Pt(18)
    purpose = left_frame.add_paragraph()
    purpose.text = str(visual.get("serves_purpose") or "")
    purpose.font.size = Pt(13)

    right = slide.shapes.add_textbox(Inches(5.1), Inches(1.15), Inches(4.0), Inches(2.65))
    right_frame = right.text_frame
    right_frame.clear()
    right_frame.paragraphs[0].text = "数学表达"
    right_frame.paragraphs[0].font.size = Pt(16)
    right_frame.paragraphs[0].font.bold = True
    for assertion in page.get("math_assertions") or []:
        if not isinstance(assertion, dict):
            continue
        paragraph = right_frame.add_paragraph()
        paragraph.text = str(assertion.get("content") or "")
        paragraph.font.size = Pt(18)

    footer = slide.shapes.add_textbox(Inches(0.75), Inches(4.25), Inches(8.4), Inches(0.8))
    footer_frame = footer.text_frame
    footer_frame.clear()
    footer_frame.paragraphs[0].text = str(page.get("evidence_requirement") or page.get("student_action") or "")
    footer_frame.paragraphs[0].font.size = Pt(15)

    return _write_notes(slide, page, sources)


def _write_notes(slide, page: dict[str, Any], sources: dict[str, Any]) -> int:
    notes = str(page.get("accuracy_notes") or "").strip()
    if not notes:
        return 0
    try:
        text_frame = slide.notes_slide.notes_text_frame
        text_frame.text = notes
        return 1
    except Exception:
        return 0


def _add_video_slide(presentation: Presentation, video_path: Path) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(8.5), Inches(0.5))
    title = title_box.text_frame.paragraphs[0]
    title.text = "导入视频"
    title.font.size = Pt(24)
    title.font.bold = True
    try:
        slide.shapes.add_movie(
            str(video_path),
            Inches(0.75),
            Inches(1.0),
            Inches(8.5),
            Inches(4.8),
            mime_type="video/mp4",
        )
    except Exception as exc:
        raise PptExportError(f"视频嵌入失败：{exc}") from exc


def _source_pages(sources: dict[str, Any]) -> list[dict[str, Any]]:
    script = sources.get("ppt_page_script") if isinstance(sources.get("ppt_page_script"), dict) else {}
    pages = script.get("pages") if isinstance(script, dict) else []
    if not isinstance(pages, list):
        return []
    return [page for page in pages if isinstance(page, dict)]


def _media_count(sources: dict[str, Any]) -> int:
    visual_asset = sources.get("ppt_visual_asset") if isinstance(sources.get("ppt_visual_asset"), dict) else {}
    assets = visual_asset.get("assets") if isinstance(visual_asset, dict) else []
    if not isinstance(assets, list):
        return 0
    return len([asset for asset in assets if isinstance(asset, dict) and asset.get("status") in {"approved", "generated"}])


def _maybe_video_path(project_dir: Path, project_config: Any, allow_placeholder: bool) -> Path | None:
    config = project_config if isinstance(project_config, dict) else {}
    embed_video_path = config.get("_embed_video_path")
    if isinstance(embed_video_path, str) and embed_video_path:
        video_path = (project_dir / embed_video_path).resolve()
        try:
            video_path.relative_to(project_dir.resolve())
        except ValueError as exc:
            raise PptExportError("视频路径不能指向项目目录外") from exc
        if video_path.exists():
            return video_path
    if not config.get("embed_video_in_ppt"):
        return None
    video_path = project_dir / FINAL_VIDEO_REL_PATH
    if not video_path.exists():
        if not allow_placeholder:
            raise PptExportError("真实视频模式下缺少已合成的 final_video.mp4")
        video_path = ensure_final_video_output(project_dir)
    if not allow_placeholder and video_path.read_bytes() == PLACEHOLDER_MP4:
        raise PptExportError("真实视频模式下 final_video.mp4 不能是占位视频")
    return video_path


def _safe_filename(value: str) -> str:
    safe = "".join(ch if ch.isalnum() or ch in {"-", "_"} else "-" for ch in value.strip())
    safe = "-".join(part for part in safe.split("-") if part)
    return safe[:80] or "shan-hai"
