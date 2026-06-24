import operator
import re
import sqlite3
from pathlib import Path
from typing import Any

from pptx import Presentation

from .control_plane import ControlPlaneStore
from .store import ProjectStore


class RuleExecutionError(ValueError):
    def __init__(self, result: dict[str, Any]):
        super().__init__(result["message"])
        self.result = result

    @property
    def rule_id(self) -> str:
        return str(self.result["rule_id"])

    @property
    def code(self) -> str:
        return f"RULE_VIOLATION_{self.rule_id}"

    @property
    def details(self) -> dict[str, Any]:
        nested_details = self.result.get("details")
        if isinstance(nested_details, dict):
            return {**self.result, **nested_details, "details": nested_details}
        return self.result


class RuleHardBlockError(RuleExecutionError):
    pass


class RuleWarningError(ValueError):
    def __init__(self, warnings: list[dict[str, Any]]):
        super().__init__("规则 warning 需要用户确认后才能继续")
        self.warnings = warnings

    @property
    def details(self) -> dict[str, Any]:
        return {"warnings": self.warnings}


class RuleExecutor:
    def __init__(self, control_plane: ControlPlaneStore | None):
        self.control_plane = control_plane

    def run_for_event(
        self,
        conn: sqlite3.Connection,
        store: ProjectStore,
        project_id: str,
        node_id: str,
        trigger_event: str,
        content: dict[str, Any] | None,
        *,
        version_id: str | None = None,
        override_warning_rule_ids: list[str] | None = None,
        override_reason: str | None = None,
    ) -> list[dict[str, Any]]:
        override_ids = set(override_warning_rule_ids or [])
        event_rules = self._rules_for(project_id, node_id, trigger_event)
        results: list[dict[str, Any]] = []
        warnings: list[dict[str, Any]] = []
        for severity in ["hard_block", "warning", "info"]:
            for rule in [item for item in event_rules if item.get("severity") == severity]:
                check = self._run_check(rule, content or {}, conn, store, project_id)
                if check is None:
                    self._record_unimplemented_rule(store, conn, project_id, node_id, version_id, rule, trigger_event)
                    continue
                passed, details = check
                if severity == "warning" and not passed and rule["rule_id"] in override_ids:
                    details = {**details, "override": True, "override_reason": override_reason}
                    results.append(self._record(store, conn, project_id, node_id, version_id, rule, trigger_event, True, details))
                    continue
                result = self._record(store, conn, project_id, node_id, version_id, rule, trigger_event, passed, details)
                results.append(result)
                if severity == "hard_block" and not passed:
                    conn.commit()
                    raise RuleHardBlockError(result)
                if severity == "warning" and not passed:
                    warnings.append(result)
        if warnings:
            conn.commit()
            raise RuleWarningError(warnings)
        return results

    def coverage(self) -> dict[str, Any]:
        rules = self._all_rules()
        rows = []
        for rule in rules:
            implemented = self._is_implemented(rule)
            handled_by = self._handled_by(rule)
            rows.append(
                {
                    "rule_id": rule["rule_id"],
                    "title": rule.get("title"),
                    "trigger_node": rule.get("trigger_node"),
                    "trigger_event": rule.get("trigger_event"),
                    "severity": rule.get("severity"),
                    "declared_executor": rule.get("executor"),
                    "executor_type": handled_by or ("builtin" if implemented else "unimplemented"),
                    "handled_by": handled_by or "RuleExecutor",
                    "implemented": implemented,
                    "enabled": bool(rule.get("enabled", True)),
                    "version_id": rule.get("version_id"),
                }
            )
        unimplemented = [rule for rule in rows if not rule["implemented"]]
        return {
            "loaded_count": len(rows),
            "implemented_count": len(rows) - len(unimplemented),
            "unimplemented_count": len(unimplemented),
            "rules": rows,
            "implemented_rule_ids": [rule["rule_id"] for rule in rows if rule["implemented"]],
            "unimplemented_rule_ids": [rule["rule_id"] for rule in unimplemented],
            "unimplemented_hard_block_rule_ids": [rule["rule_id"] for rule in unimplemented if rule.get("severity") == "hard_block"],
        }

    def summary_for_node(self, node_id: str) -> dict[str, Any]:
        applicable = [rule for rule in self._all_rules() if self._node_matches(rule.get("trigger_node"), node_id)]
        unimplemented_hard_blocks = sorted(
            {
                str(rule["rule_id"])
                for rule in applicable
                if rule.get("severity") == "hard_block" and not self._is_implemented(rule)
            }
        )
        return {
            "unimplemented_hard_block_count": len(unimplemented_hard_blocks),
            "unimplemented_hard_block_rule_ids": unimplemented_hard_blocks,
        }

    def _record_unimplemented_rule(
        self,
        store: ProjectStore,
        conn: sqlite3.Connection,
        project_id: str,
        node_id: str,
        version_id: str | None,
        rule: dict[str, Any],
        trigger_event: str,
    ) -> dict[str, Any] | None:
        if rule.get("severity") != "hard_block":
            return None
        return self._record(
            store,
            conn,
            project_id,
            node_id,
            version_id,
            rule,
            trigger_event,
            True,
            {"implemented": False, "executor_type": "unimplemented", "coverage_visible": True},
        )

    def _record(
        self,
        store: ProjectStore,
        conn: sqlite3.Connection,
        project_id: str,
        node_id: str,
        version_id: str | None,
        rule: dict[str, Any],
        trigger_event: str,
        passed: bool,
        details: dict[str, Any],
    ) -> dict[str, Any]:
        message = rule.get("action_message") or rule.get("title") or rule["rule_id"]
        record = store.record_rule_result(
            conn,
            project_id=project_id,
            node_id=node_id,
            version_id=version_id,
            rule_id=rule["rule_id"],
            trigger_event=trigger_event,
            severity=rule["severity"],
            passed=passed,
            message=message,
            details={
                "rule_id": rule["rule_id"],
                "rule_version_id": rule.get("version_id"),
                "trigger_event": trigger_event,
                "severity": rule["severity"],
                **details,
            },
        )
        return {
            "rule_id": rule["rule_id"],
            "trigger_event": trigger_event,
            "severity": rule["severity"],
            "passed": passed,
            "message": message,
            "details": record["details"],
        }

    def _run_check(
        self,
        rule: dict[str, Any],
        content: dict[str, Any],
        conn: sqlite3.Connection,
        store: ProjectStore,
        project_id: str,
    ) -> tuple[bool, dict[str, Any]] | None:
        if rule.get("enabled") is False:
            return True, {"disabled": True}
        check = rule.get("check") if isinstance(rule.get("check"), dict) else {}
        check_type = check.get("type")
        if check_type == "all":
            return self._check_all(check, content, conn, store, project_id)
        if check_type == "field_exists":
            value = _field_value(content, str(check.get("field") or ""))
            return _is_present(value), {"field": check.get("field"), "value_present": _is_present(value)}
        if check_type == "field_equals":
            actual = _field_value(content, str(check.get("field") or ""))
            return actual == check.get("value"), {"field": check.get("field"), "actual": actual, "expected": check.get("value")}
        if check_type == "field_compare":
            return _check_field_compare(content, check)
        if check_type == "enum_in":
            actual = _field_value(content, str(check.get("field") or ""))
            values = list(check.get("values") or [])
            return actual in values, {"field": check.get("field"), "actual": actual, "values": values}
        if check_type == "list_contains":
            actual = _field_value(content, str(check.get("field") or ""))
            required = list(check.get("values") or [])
            actual_set = set(actual or []) if isinstance(actual, list) else set()
            return set(required).issubset(actual_set), {"field": check.get("field"), "missing": sorted(set(required) - actual_set)}
        if check_type == "list_min_length":
            actual = _field_value(content, str(check.get("field") or ""))
            length = len(actual) if isinstance(actual, list) else 0
            minimum = int(check.get("min") or 0)
            return length >= minimum, {"field": check.get("field"), "length": length, "min": minimum}
        if check_type == "list_each":
            return self._check_list_each(check, content, conn, store, project_id)
        if check_type == "nested_list_each":
            return self._check_nested_list_each(check, content, conn, store, project_id)
        if check_type == "object_count_compare":
            return _check_object_count_compare(content, check)
        if check_type == "text_not_contains":
            text = str(_field_value(content, str(check.get("field") or "")) or "")
            patterns = [str(item) for item in check.get("patterns") or []]
            hits = [pattern for pattern in patterns if re.search(pattern, text, flags=re.IGNORECASE)]
            return not hits, {"field": check.get("field"), "hits": hits}
        if check_type == "path_under":
            value = str(_field_value(content, str(check.get("field") or "")) or "")
            prefix = str(check.get("prefix") or "")
            return value.startswith(prefix), {"field": check.get("field"), "path": value, "prefix": prefix}
        if check_type in {"text_scan", "pptx_visible_text_not_contains"}:
            return self._check_pptx_visible_text(content, conn, store, project_id, check)
        if check_type in {"dependency_passable", "dependency_check"}:
            return None
        return None

    def _check_all(
        self,
        check: dict[str, Any],
        content: dict[str, Any],
        conn: sqlite3.Connection,
        store: ProjectStore,
        project_id: str,
    ) -> tuple[bool, dict[str, Any]] | None:
        child_results = []
        for child in [item for item in check.get("checks") or [] if isinstance(item, dict)]:
            result = self._run_check({"check": child, "enabled": True}, content, conn, store, project_id)
            if result is None:
                child_results.append({"implemented": False, "check": child})
                continue
            passed, details = result
            child_results.append({"passed": passed, **details})
        if not child_results:
            return None
        return all(item.get("passed") is True for item in child_results), {"checks": child_results}

    def _check_list_each(
        self,
        check: dict[str, Any],
        content: dict[str, Any],
        conn: sqlite3.Connection,
        store: ProjectStore,
        project_id: str,
    ) -> tuple[bool, dict[str, Any]]:
        field = str(check.get("field") or "")
        values = _field_value(content, field)
        item_name = str(check.get("item_name") or "item")
        if not isinstance(values, list):
            return False, {"field": field, "violations": [{"reason": "list_required", "actual_type": type(values).__name__}]}
        violations = []
        for index, item in enumerate(values):
            if not isinstance(item, dict):
                violations.append({"index": index, "reason": "item_must_be_object", "item_name": item_name})
                continue
            failed = self._failed_child_checks(check, item, conn, store, project_id)
            if failed:
                violations.append(
                    {
                        "index": index,
                        "item_name": item_name,
                        "item_id": item.get(f"{item_name}_id") or item.get("id") or item.get("character_id"),
                        "checks": failed,
                    }
                )
        return not violations, {"field": field, "violations": violations}

    def _check_nested_list_each(
        self,
        check: dict[str, Any],
        content: dict[str, Any],
        conn: sqlite3.Connection,
        store: ProjectStore,
        project_id: str,
    ) -> tuple[bool, dict[str, Any]]:
        list_field = str(check.get("list_field") or "")
        nested_field = str(check.get("nested_list_field") or "")
        parents = _field_value(content, list_field)
        item_name = str(check.get("item_name") or "item")
        if not isinstance(parents, list):
            return False, {"field": list_field, "violations": [{"reason": "list_required", "actual_type": type(parents).__name__}]}
        violations = []
        for parent_index, parent in enumerate(parents):
            if not isinstance(parent, dict):
                continue
            nested = _field_value(parent, nested_field)
            if not isinstance(nested, list):
                continue
            for item_index, item in enumerate(nested):
                if not isinstance(item, dict):
                    violations.append(
                        {
                            "parent_index": parent_index,
                            "index": item_index,
                            "reason": "item_must_be_object",
                            "item_name": item_name,
                        }
                    )
                    continue
                failed = self._failed_child_checks(check, item, conn, store, project_id)
                if failed:
                    violations.append(
                        {
                            "parent_index": parent_index,
                            "page_index": parent.get("page_index", parent_index + 1),
                            "index": item_index,
                            "item_name": item_name,
                            "checks": failed,
                        }
                    )
        return not violations, {"field": f"{list_field}.{nested_field}", "violations": violations}

    def _failed_child_checks(
        self,
        check: dict[str, Any],
        content: dict[str, Any],
        conn: sqlite3.Connection,
        store: ProjectStore,
        project_id: str,
    ) -> list[dict[str, Any]]:
        failed = []
        for child in [item for item in check.get("checks") or [] if isinstance(item, dict)]:
            result = self._run_check({"check": child, "enabled": True}, content, conn, store, project_id)
            if result is None:
                failed.append({"implemented": False, "check": child})
                continue
            passed, details = result
            if not passed:
                failed.append(details)
        return failed

    def _check_pptx_visible_text(
        self,
        content: dict[str, Any],
        conn: sqlite3.Connection,
        store: ProjectStore,
        project_id: str,
        check: dict[str, Any],
    ) -> tuple[bool, dict[str, Any]]:
        forbidden_patterns = [str(item) for item in check.get("forbidden_patterns") or check.get("patterns") or []]
        pptx_path = content.get("pptx_path") if isinstance(content.get("pptx_path"), str) else None
        if not pptx_path:
            return False, {
                "pptx_path": None,
                "error_code": "PPTX_PATH_MISSING",
                "violations": [{"source": "pptx_path", "error_code": "PPTX_PATH_MISSING", "text": ""}],
            }
        pptx_probe = self._pptx_visible_text_values(store, project_id, pptx_path)
        pptx_error = next((item for item in pptx_probe if "error_code" in item), None)
        if pptx_error:
            return False, {"pptx_path": pptx_path, "error_code": pptx_error["error_code"], "violations": [pptx_error]}
        text_values = [{"source": "content_json", "text": value} for value in RuleExecutor._collect_strings(content)]
        text_values.extend(pptx_probe)
        hits = []
        for item in text_values:
            value = item["text"]
            for pattern in forbidden_patterns:
                if re.search(pattern, value, flags=re.IGNORECASE):
                    hits.append({"source": item["source"], "pattern": pattern, "text": value[:120]})
        return not hits, {"pptx_path": pptx_path, "violations": hits}

    def _pptx_visible_text_values(self, store: ProjectStore, project_id: str, rel_path: str) -> list[dict[str, str]]:
        project_dir = Path(store.get_project(project_id)["project_dir"])
        project_root = project_dir.resolve()
        pptx_path = (project_dir / rel_path).resolve()
        if not _is_within(pptx_path, project_root):
            return [{"source": "pptx_path", "error_code": "PPTX_PATH_OUTSIDE_PROJECT", "text": rel_path}]
        if pptx_path.suffix.lower() != ".pptx":
            return [{"source": "pptx_path", "error_code": "PPTX_EXTENSION_INVALID", "text": rel_path}]
        if not pptx_path.exists():
            return [{"source": "pptx_path", "error_code": "PPTX_FILE_NOT_FOUND", "text": rel_path}]
        try:
            presentation = Presentation(str(pptx_path))
        except Exception:
            return [{"source": "pptx_path", "error_code": "PPTX_PARSE_FAILED", "text": rel_path}]
        values: list[dict[str, str]] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            for shape_index, shape in enumerate(slide.shapes, start=1):
                if not getattr(shape, "has_text_frame", False):
                    continue
                text = "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs).strip()
                if text:
                    values.append({"source": "pptx_shape_text", "text": text, "location": f"slide:{slide_index}/shape:{shape_index}"})
        return values

    @staticmethod
    def _collect_strings(value: Any) -> list[str]:
        if isinstance(value, str):
            return [value]
        if isinstance(value, dict):
            values = []
            for key, item in value.items():
                if key == "notes_text":
                    continue
                values.extend(RuleExecutor._collect_strings(item))
            return values
        if isinstance(value, list):
            values = []
            for item in value:
                values.extend(RuleExecutor._collect_strings(item))
            return values
        return []

    def _rules_for(self, project_id: str | None, node_id: str, trigger_event: str) -> list[dict[str, Any]]:
        return [
            rule
            for rule in self._rules(project_id)
            if rule.get("enabled", True)
            and rule.get("trigger_event") == trigger_event
            and self._node_matches(rule.get("trigger_node"), node_id)
        ]

    def _rule_by_id(self, project_id: str | None, rule_id: str) -> dict[str, Any] | None:
        return next((rule for rule in self._rules(project_id) if rule.get("rule_id") == rule_id), None)

    def _all_rules(self) -> list[dict[str, Any]]:
        return self._rules(None)

    def _rules(self, project_id: str | None) -> list[dict[str, Any]]:
        if self.control_plane is None:
            return []
        return self.control_plane.rules_for_project(project_id)

    def _is_implemented(self, rule: dict[str, Any]) -> bool:
        if self._handled_by(rule) == "StateEngine":
            return True
        return self._run_check(rule, {}, sqlite3.connect(":memory:"), _NullStore(), "") is not None

    @staticmethod
    def _handled_by(rule: dict[str, Any]) -> str | None:
        check = rule.get("check") if isinstance(rule.get("check"), dict) else {}
        if check.get("type") in {"dependency_passable", "dependency_check"}:
            return "StateEngine"
        return None

    @staticmethod
    def _node_matches(trigger_node: str | None, node_id: str) -> bool:
        if trigger_node in {None, "", "所有"}:
            return True
        candidates = [part.strip() for part in re.split(r"[|/]", str(trigger_node)) if part.strip()]
        return node_id in candidates


class _NullStore:
    def get_project(self, project_id: str) -> dict[str, Any]:
        raise KeyError(project_id)


def _field_value(content: dict[str, Any], field: str) -> Any:
    current: Any = content
    for part in [item for item in field.split(".") if item]:
        if not isinstance(current, dict):
            return None
        current = current.get(part)
    return current


def _is_present(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, (list, dict)):
        return bool(value)
    return True


def _check_field_compare(content: dict[str, Any], check: dict[str, Any]) -> tuple[bool, dict[str, Any]]:
    actual = _field_value(content, str(check.get("field") or ""))
    expected = check.get("value")
    op_name = str(check.get("op") or "==")
    if op_name not in COMPARISON_OPS:
        return False, {"field": check.get("field"), "actual": actual, "expected": expected, "op": op_name, "error": "unsupported_operator"}
    return _compare_values(actual, expected, op_name), {"field": check.get("field"), "actual": actual, "expected": expected, "op": op_name}


def _check_object_count_compare(content: dict[str, Any], check: dict[str, Any]) -> tuple[bool, dict[str, Any]]:
    target = _field_value(content, str(check.get("field") or ""))
    if not isinstance(target, dict):
        return False, {"field": check.get("field"), "count": 0, "matched_keys": [], "error": "object_required"}
    value_check = {
        "op": check.get("value_op") or ">=",
        "value": check.get("value", 1),
    }
    matched_keys = [
        key
        for key, value in target.items()
        if _compare_values(value, value_check["value"], str(value_check["op"]))
    ]
    count = len(matched_keys)
    expected_count = check.get("count", 0)
    count_op = str(check.get("count_op") or ">=")
    return _compare_values(count, expected_count, count_op), {
        "field": check.get("field"),
        "count": count,
        "expected_count": expected_count,
        "count_op": count_op,
        "matched_keys": matched_keys,
    }


def _compare_values(actual: Any, expected: Any, op_name: str) -> bool:
    op = COMPARISON_OPS.get(op_name)
    if op is None:
        return False
    try:
        return bool(op(actual, expected))
    except TypeError:
        return False


COMPARISON_OPS = {
    "==": operator.eq,
    "!=": operator.ne,
    ">": operator.gt,
    ">=": operator.ge,
    "<": operator.lt,
    "<=": operator.le,
}


def _is_within(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False
