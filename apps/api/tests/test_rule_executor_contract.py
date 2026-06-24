import json
import sqlite3
from pathlib import Path
from typing import Any

from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app.main import create_app


def make_client(tmp_path: Path) -> TestClient:
    app = create_app(
        {
            "storage_root": str(tmp_path / "storage"),
            "workflow_root": str(Path(__file__).resolve().parents[3] / "workflow"),
            "provider_mode": "fake",
            "video_provider_mode": "placeholder",
            "image_provider_mode": "placeholder",
            "tts_provider_mode": "placeholder",
        }
    )
    return TestClient(app)


def unwrap_ok(response):
    assert response.status_code < 400, response.text
    payload = response.json()
    assert payload["ok"] is True, payload
    return payload["data"]


def create_project(client: TestClient, name: str = "RuleExecutor 契约") -> dict[str, Any]:
    return unwrap_ok(
        client.post(
            "/projects",
            json={
                "name": name,
                "subject": "math",
                "grade": "2",
                "textbook_version": "renjiao",
                "volume": "shang",
                "lesson_type": "public",
            },
        )
    )


def write_current_version(
    client: TestClient,
    project: dict[str, Any],
    node_id: str,
    content: dict[str, Any],
    status: str = "needs_review",
) -> dict[str, Any]:
    project_id = project["project_id"]
    with client.app.state.store.connect(Path(project["project_dir"])) as conn:
        return client.app.state.store.write_version(conn, project_id, node_id, content, "fixture", "fixture", status)


TEST_DEPENDENCIES = {
    "ppt_assembly_plan": ["lesson_plan"],
    "ppt_page_script": ["ppt_assembly_plan", "character_dict", "visual_contract"],
    "pptx_artifact": ["ppt_page_script", "ppt_visual_asset"],
    "final_video": ["storyboard", "intro_video_script"],
    "final_delivery": ["pptx_artifact", "final_video"],
}


def seed_approved_upstreams(client: TestClient, project: dict[str, Any], node_id: str) -> None:
    store = client.app.state.store
    project_id = project["project_id"]
    with store.connect(Path(project["project_dir"])) as conn:
        def seed(dep_id: str) -> None:
            for upstream_id in TEST_DEPENDENCIES.get(dep_id, []):
                seed(upstream_id)
            state = store.node_state(conn, project_id, dep_id)
            if state["status"] == "approved":
                return
            result = store.write_version(conn, project_id, dep_id, {"seeded": dep_id}, "fixture", "fixture", "approved")
            store.update_current_version_status(conn, result["version_id"], "approved", approved=True)

        for dependency_id in TEST_DEPENDENCIES.get(node_id, []):
            seed(dependency_id)


def write_pptx(project: dict[str, Any], rel_path: str, visible_text: str) -> None:
    pptx_path = Path(project["project_dir"]) / rel_path
    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    box.text_frame.text = visible_text
    presentation.save(pptx_path)


def rule_rows(project: dict[str, Any], rule_id: str) -> list[sqlite3.Row]:
    with sqlite3.connect(Path(project["project_dir"]) / "project.db") as conn:
        conn.row_factory = sqlite3.Row
        return conn.execute(
            "SELECT * FROM rule_result_log WHERE rule_id = ? ORDER BY created_at",
            (rule_id,),
        ).fetchall()


def transition_rows(project: dict[str, Any], node_id: str) -> list[sqlite3.Row]:
    with sqlite3.connect(Path(project["project_dir"]) / "project.db") as conn:
        conn.row_factory = sqlite3.Row
        return conn.execute(
            "SELECT * FROM state_transition_log WHERE node_id = ? ORDER BY triggered_at, rowid",
            (node_id,),
        ).fetchall()


def error_payload(response, status_code: int, code: str) -> dict[str, Any]:
    assert response.status_code == status_code, response.text
    payload = response.json()
    assert payload["ok"] is False, payload
    assert payload["error"]["code"] == code, payload
    return payload["error"]


def test_r010_blocks_downstream_generate_and_logs_state_engine_diagnostic(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]

    response = client.post(f"/projects/{project_id}/nodes/lesson_plan/generate", json={})

    error_payload(response, 409, "UPSTREAM_NOT_APPROVED")
    assert rule_rows(project, "R010") == []
    rows = transition_rows(project, "lesson_plan")
    blocked_rows = [row for row in rows if row["trigger"] == "dependency_gate_blocked"]
    assert len(blocked_rows) == 1
    reason = json.loads(blocked_rows[0]["reason"])
    assert reason["rule_id"] == "R010"
    assert reason["handled_by"] == "StateEngine"
    assert reason["blocked_dependencies"][0]["node_id"] == "textbook_parse"


def test_rule_executor_coverage_marks_yaml_rules_without_runtime_checks_unimplemented(tmp_path: Path):
    client = make_client(tmp_path)

    coverage = client.app.state.service.rule_executor.coverage()

    by_id = {rule["rule_id"]: rule for rule in coverage["rules"]}
    assert by_id["R010"]["implemented"] is True
    assert by_id["R010"]["handled_by"] == "StateEngine"
    assert by_id["R010"]["executor_type"] == "StateEngine"
    assert by_id["R001"]["implemented"] is True
    assert by_id["R004"]["implemented"] is True
    assert by_id["R005"]["implemented"] is True
    assert by_id["R006"]["implemented"] is True
    assert by_id["R026"]["implemented"] is True
    assert by_id["R030"]["implemented"] is True
    assert by_id["R011"]["implemented"] is False
    assert by_id["R011"]["executor_type"] == "unimplemented"
    assert "R011" in coverage["unimplemented_rule_ids"]
    assert "R011" in coverage["unimplemented_hard_block_rule_ids"]


def test_manifest_rule_summary_exposes_unimplemented_hard_block_rules_for_node(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]

    manifest = unwrap_ok(client.get(f"/projects/{project_id}/manifest"))
    ppt_asset = {node["node_id"]: node for node in manifest["nodes"]}["ppt_visual_asset"]

    assert "R011" in ppt_asset["rule_summary"]["unimplemented_hard_block_rule_ids"]
    assert ppt_asset["rule_summary"]["unimplemented_hard_block_count"] >= 1


def test_r004_blocks_character_dict_save_and_logs_result(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    broken = {
        "characters": [
            {
                "character_id": "real_child",
                "name": "真实儿童",
                "style_constraint": "photorealistic",
                "banned_keywords": [],
            }
        ]
    }

    response = client.post(f"/projects/{project_id}/nodes/character_dict/edit", json={"content": broken})

    error = error_payload(response, 400, "RULE_VIOLATION_R004")
    assert error["details"]["rule_id"] == "R004"
    rows = rule_rows(project, "R004")
    assert len(rows) == 1
    assert rows[0]["passed"] == 0

    transitions = transition_rows(project, "character_dict")
    assert transitions[-1]["to_status"] == "blocked"
    assert transitions[-1]["trigger"] == "hard_block_rule_hit"


def test_r006_blocks_ppt_page_script_approve(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "ppt_page_script")
    broken = {
        "pages": [
            {
                "page_index": 1,
                "math_assertions": [
                    {"content": "286 + 147 = 433", "answer": "433", "editable_layer": "ppt_image"}
                ],
            }
        ]
    }
    write_current_version(client, project, "ppt_page_script", broken)

    response = client.post(f"/projects/{project_id}/nodes/ppt_page_script/approve", json={})

    error = error_payload(response, 409, "RULE_VIOLATION_R006")
    assert "editable" in json.dumps(error["details"], ensure_ascii=False)
    rows = rule_rows(project, "R006")
    assert len(rows) == 1
    assert rows[0]["passed"] == 0


def test_r005_blocks_candidate_final_delivery_approve(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "final_delivery")
    content = {
        "lesson_plan_path": "exports/final_delivery/lesson_plan.md",
        "pptx_final_path": "exports/final_delivery/deck.pptx",
        "video_final_path": "exports/final_delivery/final_video.mp4",
        "delivery_manifest_path": "exports/final_delivery/delivery_manifest.json",
        "gate_result_json_path": "exports/final_delivery/gate_result.json",
        "gate_result_json": {"mode": "candidate", "gate_passed": False},
        "gate_passed": False,
        "qa_records": ["candidate_gate"],
        "time_stats_md_path": "exports/final_delivery/time_stats.md",
        "feedback_trigger_at": "2026-06-22T00:00:00Z",
        "source_versions": {"lesson_plan": "ver_lesson", "pptx_artifact": "ver_ppt", "final_video": "ver_video"},
        "generated_at": "2026-06-22T00:00:00Z",
    }
    write_current_version(client, project, "final_delivery", content)

    response = client.post(f"/projects/{project_id}/nodes/final_delivery/approve", json={})

    error = error_payload(response, 409, "RULE_VIOLATION_R005")
    assert error["details"]["rule_id"] == "R005"
    rows = rule_rows(project, "R005")
    assert len(rows) == 1
    assert rows[0]["passed"] == 0


def test_r030_blocks_pptx_artifact_without_visual_assets(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "pptx_artifact")
    write_pptx(project, "exports/no-visual.pptx", "学生可见层干净")
    broken = {
        "pptx_path": "exports/no-visual.pptx",
        "slide_count": 2,
        "notes_count": 2,
        "media_count": 0,
        "svg_quality_passed": True,
    }
    write_current_version(client, project, "pptx_artifact", broken)

    response = client.post(f"/projects/{project_id}/nodes/pptx_artifact/approve", json={})

    error_payload(response, 409, "RULE_VIOLATION_R030")
    rows = rule_rows(project, "R030")
    assert len(rows) == 1
    assert rows[0]["passed"] == 0


def test_r026_blocks_internal_text_in_actual_pptx_visible_shapes(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "pptx_artifact")
    write_pptx(project, "exports/internal-visible.pptx", "这里是学生可见层 QA 检查点")
    clean_json = {
        "pptx_path": "exports/internal-visible.pptx",
        "slide_count": 1,
        "notes_count": 0,
        "media_count": 1,
        "svg_quality_passed": True,
    }
    write_current_version(client, project, "pptx_artifact", clean_json)

    response = client.post(f"/projects/{project_id}/nodes/pptx_artifact/approve", json={})

    error_payload(response, 409, "RULE_VIOLATION_R026")
    rows = rule_rows(project, "R026")
    assert len(rows) == 1
    assert rows[0]["passed"] == 0
    details = json.loads(rows[0]["details_json"])
    assert details["pptx_path"] == "exports/internal-visible.pptx"
    assert details["violations"][0]["source"] == "pptx_shape_text"


def test_r026_ignores_notes_text_when_actual_pptx_visible_shapes_are_clean(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "pptx_artifact")
    write_pptx(project, "exports/clean-visible.pptx", "学生可见层干净")
    clean_visible = {
        "pptx_path": "exports/clean-visible.pptx",
        "slide_count": 1,
        "notes_count": 1,
        "media_count": 1,
        "svg_quality_passed": True,
        "notes_text": "教师备注：QA 检查点仅供内部核验",
    }
    write_current_version(client, project, "pptx_artifact", clean_visible)

    approved = unwrap_ok(client.post(f"/projects/{project_id}/nodes/pptx_artifact/approve", json={}))

    assert approved == {"node_id": "pptx_artifact", "status": "approved"}
    rows = rule_rows(project, "R026")
    assert len(rows) == 1
    assert rows[0]["passed"] == 1


def test_r026_blocks_pptx_artifact_without_pptx_path(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "pptx_artifact")
    content = {
        "slide_count": 1,
        "notes_count": 0,
        "media_count": 1,
        "svg_quality_passed": True,
    }
    write_current_version(client, project, "pptx_artifact", content)

    response = client.post(f"/projects/{project_id}/nodes/pptx_artifact/approve", json={})

    error = error_payload(response, 409, "RULE_VIOLATION_R026")
    assert error["details"]["error_code"] == "PPTX_PATH_MISSING"
    rows = rule_rows(project, "R026")
    assert len(rows) == 1
    assert rows[0]["passed"] == 0


def test_r026_blocks_pptx_artifact_when_file_is_missing(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "pptx_artifact")
    content = {
        "pptx_path": "exports/missing.pptx",
        "slide_count": 1,
        "notes_count": 0,
        "media_count": 1,
        "svg_quality_passed": True,
    }
    write_current_version(client, project, "pptx_artifact", content)

    response = client.post(f"/projects/{project_id}/nodes/pptx_artifact/approve", json={})

    error = error_payload(response, 409, "RULE_VIOLATION_R026")
    assert error["details"]["error_code"] == "PPTX_FILE_NOT_FOUND"
    rows = rule_rows(project, "R026")
    assert len(rows) == 1
    assert rows[0]["passed"] == 0


def test_r026_blocks_pptx_artifact_when_path_is_outside_project(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "pptx_artifact")
    content = {
        "pptx_path": "../outside.pptx",
        "slide_count": 1,
        "notes_count": 0,
        "media_count": 1,
        "svg_quality_passed": True,
    }
    write_current_version(client, project, "pptx_artifact", content)

    response = client.post(f"/projects/{project_id}/nodes/pptx_artifact/approve", json={})

    error = error_payload(response, 409, "RULE_VIOLATION_R026")
    assert error["details"]["error_code"] == "PPTX_PATH_OUTSIDE_PROJECT"


def test_r026_blocks_pptx_artifact_when_path_is_not_pptx(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "pptx_artifact")
    bad_path = Path(project["project_dir"]) / "exports" / "artifact.txt"
    bad_path.parent.mkdir(parents=True, exist_ok=True)
    bad_path.write_text("not a pptx", encoding="utf-8")
    content = {
        "pptx_path": "exports/artifact.txt",
        "slide_count": 1,
        "notes_count": 0,
        "media_count": 1,
        "svg_quality_passed": True,
    }
    write_current_version(client, project, "pptx_artifact", content)

    response = client.post(f"/projects/{project_id}/nodes/pptx_artifact/approve", json={})

    error = error_payload(response, 409, "RULE_VIOLATION_R026")
    assert error["details"]["error_code"] == "PPTX_EXTENSION_INVALID"


def test_r026_blocks_pptx_artifact_when_pptx_cannot_be_parsed(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "pptx_artifact")
    bad_path = Path(project["project_dir"]) / "exports" / "broken.pptx"
    bad_path.parent.mkdir(parents=True, exist_ok=True)
    bad_path.write_text("not a real pptx", encoding="utf-8")
    content = {
        "pptx_path": "exports/broken.pptx",
        "slide_count": 1,
        "notes_count": 0,
        "media_count": 1,
        "svg_quality_passed": True,
    }
    write_current_version(client, project, "pptx_artifact", content)

    response = client.post(f"/projects/{project_id}/nodes/pptx_artifact/approve", json={})

    error = error_payload(response, 409, "RULE_VIOLATION_R026")
    assert error["details"]["error_code"] == "PPTX_PARSE_FAILED"


def test_r001_warns_but_allows_override_for_non_male_voice(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "final_video")
    content = {
        "clip_count": 1,
        "clips": [],
        "model_audio_policy": "verified_chinese",
        "english_audio_detected": False,
        "voice_gender": "female",
        "voice_language": "zh-CN",
        "audio_verified": True,
        "narration_audio_path": "audio/narration.mp3",
    }
    warning_content = {**content, "audio_verified": False, "narration_audio_path": ""}
    write_current_version(client, project, "final_video", warning_content)
    warning = client.post(f"/projects/{project_id}/nodes/final_video/approve", json={})
    error = error_payload(warning, 409, "RULE_WARNING")
    assert error["details"]["warnings"][0]["rule_id"] == "R001"
    assert error["details"]["warnings"][0]["severity"] == "warning"

    approved = unwrap_ok(
        client.post(
            f"/projects/{project_id}/nodes/final_video/approve",
            json={"override_warning_rule_ids": ["R001"], "override_reason": "教师选择女声旁白"},
        )
    )

    assert approved == {"node_id": "final_video", "status": "approved"}
    rows = rule_rows(project, "R001")
    assert [row["passed"] for row in rows] == [0, 1]
    details = json.loads(rows[-1]["details_json"])
    assert details["override"] is True
    assert details["override_reason"] == "教师选择女声旁白"


def test_warning_requires_override_and_override_is_logged(tmp_path: Path):
    client = make_client(tmp_path)
    project = create_project(client)
    project_id = project["project_id"]
    seed_approved_upstreams(client, project, "ppt_assembly_plan")
    warning_content = {
        "page_count_target": 6,
        "page_type_quota": {
            "life_observation": 1,
            "role_task": 1,
            "inquiry_operation": 1,
            "step_reveal": 1,
            "practice_challenge": 1,
            "blackboard_summary": 0,
        },
    }
    write_current_version(client, project, "ppt_assembly_plan", warning_content)

    warning = client.post(f"/projects/{project_id}/nodes/ppt_assembly_plan/approve", json={})
    error = error_payload(warning, 409, "RULE_WARNING")
    assert error["details"]["warnings"][0]["rule_id"] == "R023"

    approved = unwrap_ok(
        client.post(
            f"/projects/{project_id}/nodes/ppt_assembly_plan/approve",
            json={"override_warning_rule_ids": ["R023"], "override_reason": "内测页数压缩"},
        )
    )

    assert approved == {"node_id": "ppt_assembly_plan", "status": "approved"}
    rows = rule_rows(project, "R023")
    assert [row["passed"] for row in rows] == [0, 1]
    details = json.loads(rows[-1]["details_json"])
    assert details["override"] is True
    assert details["override_reason"] == "内测页数压缩"
